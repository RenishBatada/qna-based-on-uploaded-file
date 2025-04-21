import os
from dotenv import load_dotenv
load_dotenv()

from langchain.text_splitter import RecursiveCharacterTextSplitter
from langchain_ollama import OllamaEmbeddings
from langchain_pinecone import PineconeVectorStore
from langchain_core.documents import Document
import pandas as pd
import docx
from pptx import Presentation
import tempfile
import io
import PyPDF2

# Initialize embeddings
embedding = OllamaEmbeddings(
    model="llama3.2:latest",
    base_url="http://127.0.0.1:11434",
)

from langchain_google_genai import GoogleGenerativeAIEmbeddings
embedding = GoogleGenerativeAIEmbeddings(model="models/text-embedding-004")


def read_docx_text(file_path):
    """Extract text from a DOCX file, returning a list of paragraph texts."""
    doc = docx.Document(file_path)
    return [para.text for para in doc.paragraphs if para.text.strip()]

def read_pptx_text(file_path):
    """Extract text from a PPTX file, returning a list of slide texts."""
    prs = Presentation(file_path)
    slide_texts = []
    for slide_num, slide in enumerate(prs.slides):
        slide_text = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                slide_text.append(shape.text)
        if slide_text:
            slide_texts.append(f"Slide {slide_num+1}: " + "\n".join(slide_text))
    return slide_texts

def read_pdf_text(file_obj):
    """Extract text from a PDF file, returning a list of page texts and page numbers."""
    reader = PyPDF2.PdfReader(file_obj)
    page_texts = []
    page_numbers = []
    
    for i, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        if text.strip():
            page_texts.append(text)
            page_numbers.append(i + 1)
    
    return page_texts, page_numbers

def process_file_content(file, file_name, file_type):
    """
    Process uploaded file content based on file type.
    
    Args:
        file: The file object
        file_name: Name of the file
        file_type: Type of the file
        
    Returns:
        Tuple of (file_content, page_texts, page_numbers)
    """
    file_content = ""
    page_texts = []
    page_numbers = []
    
    if file_name.endswith((".csv", ".xls", ".xlsx")):
        if file_name.endswith(".csv"):
            df = pd.read_csv(file)
        else:
            df = pd.read_excel(file)
        return df, None, None
        
    elif file_name.endswith(".pdf"):
        page_texts, page_numbers = read_pdf_text(file)
        file_content = "\n\n".join(page_texts)
        
    elif file_name.endswith(".txt"):
        content = file.read().decode("utf-8")
        page_texts = [content]
        page_numbers = [1]
        file_content = content
        
    elif file_name.endswith(".docx"):
        with tempfile.NamedTemporaryFile(delete=False, suffix=".docx") as tmp:
            tmp.write(file.read())
            tmp_path = tmp.name
        page_texts = read_docx_text(tmp_path)
        page_numbers = list(range(1, len(page_texts) + 1))
        file_content = "\n\n".join(page_texts)
        os.unlink(tmp_path)
        
    elif file_name.endswith(".pptx"):
        with tempfile.NamedTemporaryFile(delete=False, suffix=".pptx") as tmp:
            tmp.write(file.read())
            tmp_path = tmp.name
        page_texts = read_pptx_text(tmp_path)
        page_numbers = list(range(1, len(page_texts) + 1))
        file_content = "\n\n".join(page_texts)
        os.unlink(tmp_path)
    
    return file_content, page_texts, page_numbers

def ingest_text_chunks(texts: list[str], file_name: str, file_type: str, page_numbers: list[int], metadata: dict = None):
    """
    Ingest text chunks into the vector database with appropriate metadata.
    
    Args:
        texts: List of text chunks (e.g., pages from a document)
        file_name: Name of the file
        file_type: Type of the file
        page_numbers: List of page numbers corresponding to each text chunk
        metadata: Optional additional metadata to include
    """
    print(f"Ingesting {len(texts)} chunks for {file_name}")
    
    # Initialize text splitter for chunking
    text_splitter = RecursiveCharacterTextSplitter(
        chunk_size=2000,
        chunk_overlap=200,
        length_function=len,
    )
    
    # Create documents with metadata
    documents = []
    for i, (text, page_num) in enumerate(zip(texts, page_numbers)):
        # Split text into smaller chunks if needed
        chunks = text_splitter.split_text(text)
        
        # Create documents for each chunk
        for chunk in chunks:
            # Combine base metadata with additional metadata if provided
            doc_metadata = {
                "source": file_name,
                "file_type": file_type,
                "page": str(page_num)
            }
            if metadata:
                doc_metadata.update(metadata)
            
            documents.append(
                Document(
                    page_content=chunk,
                    metadata=doc_metadata
                )
            )
    
    print(f"Created {len(documents)} documents")
    
    # Get vector store
    Index_name = os.getenv("INDEX_NAME")
    vector_store = PineconeVectorStore(index_name=Index_name, embedding=embedding)
    
    # Add documents to vector store
    vector_store.add_documents(documents)
    print(f"Added {len(documents)} documents to vector store")

def ingest_tabular_chunks(df: pd.DataFrame, file_name: str, file_type: str):
    """
    Ingest tabular data (CSV, Excel) into the vector database, chunking by multiple rows to reach ~2000 characters per chunk.
    Args:
        df: Pandas DataFrame containing the tabular data
        file_name: Name of the file
        file_type: Type of the file
    """
    print(f"Ingesting tabular data into vectorstore for {file_name}...")

    all_documents = []
    chunk_rows = []
    chunk_char_count = 0
    chunk_size = 2000  # Target chunk size in characters
    start_row = 0

    # Convert DataFrame to CSV header
    header = ','.join(df.columns) + '\n'
    for i, row in df.iterrows():
        row_csv = ','.join([str(row[col]) for col in df.columns]) + '\n'
        if chunk_char_count + len(row_csv) > chunk_size and chunk_rows:
            # Store current chunk
            chunk_csv = header + ''.join(chunk_rows)
            all_documents.append(
                Document(
                    page_content=chunk_csv,
                    metadata={
                        "source": file_name,
                        "page": f"{start_row+1}-{i}",
                        "file_type": file_type,
                        "row_start": start_row,
                        "row_end": i-1
                    }
                )
            )
            chunk_rows = []
            chunk_char_count = 0
            start_row = i
        chunk_rows.append(row_csv)
        chunk_char_count += len(row_csv)

    # Store any remaining rows as last chunk
    if chunk_rows:
        chunk_csv = header + ''.join(chunk_rows)
        all_documents.append(
            Document(
                page_content=chunk_csv,
                metadata={
                    "source": file_name,
                    "page": f"{start_row+1}-{start_row+len(chunk_rows)}",
                    "file_type": file_type,
                    "row_start": start_row,
                    "row_end": start_row+len(chunk_rows)-1
                }
            )
        )

    print(f"Total records/chunks created: {len(all_documents)}")
    if all_documents:
        PineconeVectorStore.from_documents(
            all_documents, embedding, index_name=os.getenv("INDEX_NAME")
        )
    print(" Tabular ingestion complete.")

def clear_vector_store():
    """Clear all documents from the vector store."""
    try:
        vector_store = PineconeVectorStore(
            index_name=os.getenv("INDEX_NAME"),
            embedding=embedding
        )
        vector_store.delete(delete_all=True)
        print(" Vector store cleared successfully.")
        return True
    except Exception as e:
        print(f"Error clearing vector store: {e}")
        return False
